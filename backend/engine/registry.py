from __future__ import annotations

import csv
import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .extraction import TYPE_ALIASES, normalize
from .models import EquipmentCandidate, EquipmentType, LocalExtraction
from .office import to_xlsx

NUMBER_RE = re.compile(r"\b\d{5,8}\b")


@dataclass(frozen=True)
class RegistryLink:
    reference_number: str
    isolation_type: EquipmentType
    isolation_number: str
    municipality: str = ""


REFERENCE_HEADERS = {
    "EQUIPAMENTO REFERENCIA",
    "NUMERO REFERENCIA",
    "REFERENCE NUMBER",
    "EQUIPAMENTO JUSANTE",
    "JUSANTE",
    "CHILD NUMBER",
    "TRANSFORMADOR",
}
ISOLATION_NUMBER_HEADERS = {
    "NUMERO ISOLAMENTO",
    "EQUIPAMENTO ISOLAMENTO",
    "DISPOSITIVO ISOLAMENTO",
    "EQUIPAMENTO MONTANTE",
    "MONTANTE",
    "UPSTREAM NUMBER",
    "ISOLAMENTO",
}
ISOLATION_TYPE_HEADERS = {
    "TIPO ISOLAMENTO",
    "TIPO MONTANTE",
    "UPSTREAM TYPE",
    "DISPOSITIVO TIPO",
}
MUNICIPALITY_HEADERS = {"MUNICIPIO", "CIDADE", "MUNICIPALITY"}


def _first(record: dict[str, str], aliases: set[str]) -> str:
    return next((record[key] for key in aliases if record.get(key)), "")


def _number(value: str) -> str:
    match = NUMBER_RE.search(str(value))
    return match.group(0) if match else ""


def _type(value: str) -> EquipmentType | None:
    normalized = normalize(str(value))
    if normalized in TYPE_ALIASES:
        return TYPE_ALIASES[normalized]
    token = normalized.split(" ", 1)[0]
    try:
        return EquipmentType(token)
    except ValueError:
        return None


def _csv_rows(path: Path) -> Iterable[list[str]]:
    raw = path.read_bytes()
    text = next(
        (raw.decode(encoding) for encoding in ("utf-8-sig", "latin-1") if _decodable(raw, encoding)),
        raw.decode("utf-8", errors="replace"),
    )
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    yield from csv.reader(text.splitlines(), dialect=dialect)


def _decodable(raw: bytes, encoding: str) -> bool:
    try:
        raw.decode(encoding)
        return True
    except UnicodeDecodeError:
        return False


def _spreadsheet_rows(path: Path) -> Iterable[list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(values_only=True):
            yield ["" if value is None else str(value) for value in row]
    finally:
        workbook.close()


def load_registry(path: Path, *, work_dir: Path, libreoffice_bin: str) -> list[RegistryLink]:
    source = path
    if path.suffix.lower() == ".xls":
        source = to_xlsx(path, work_dir, libreoffice_bin)
    rows = list(_csv_rows(source) if source.suffix.lower() == ".csv" else _spreadsheet_rows(source))
    if len(rows) < 2:
        return []
    headers = [normalize(value) for value in rows[0]]
    links: list[RegistryLink] = []
    for values in rows[1:]:
        record = {
            header: str(values[index]).strip()
            for index, header in enumerate(headers)
            if header and index < len(values)
        }
        reference = _number(_first(record, REFERENCE_HEADERS))
        isolation = _number(_first(record, ISOLATION_NUMBER_HEADERS))
        isolation_type = _type(_first(record, ISOLATION_TYPE_HEADERS))
        if not reference or not isolation or isolation_type is None:
            continue
        links.append(
            RegistryLink(
                reference_number=reference,
                isolation_type=isolation_type,
                isolation_number=isolation,
                municipality=_first(record, MUNICIPALITY_HEADERS),
            )
        )
    return links


def enrich_from_registry(extraction: LocalExtraction, links: list[RegistryLink]) -> LocalExtraction:
    project_numbers = set(extraction.identifiers)
    project_numbers.update(action.number for action in extraction.actions)
    project_numbers.update(candidate.number for candidate in extraction.candidates)
    municipality = normalize(extraction.metadata.municipio)
    for link in links:
        if link.reference_number not in project_numbers:
            continue
        if link.municipality and municipality and normalize(link.municipality) != municipality:
            continue
        evidence = f"cadastro: isolamento a montante de {link.reference_number}"
        existing = next(
            (
                candidate
                for candidate in extraction.candidates
                if candidate.number == link.isolation_number
                and candidate.equipment_type == link.isolation_type
            ),
            None,
        )
        if existing is not None:
            existing.score = max(existing.score, 0.96)
            if evidence not in existing.evidence:
                existing.evidence.append(evidence)
        else:
            extraction.candidates.append(
                EquipmentCandidate(
                    equipment_type=link.isolation_type,
                    number=link.isolation_number,
                    score=0.96,
                    evidence=[evidence],
                )
            )
        extraction.registry_identifiers.append(link.isolation_number)
    extraction.registry_identifiers = sorted(set(extraction.registry_identifiers))
    extraction.candidates.sort(key=lambda candidate: (-candidate.score, candidate.number))
    return extraction
